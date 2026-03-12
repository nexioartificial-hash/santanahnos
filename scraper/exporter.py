"""
Exportación a XLSX, CSV y JSON para el scraper de repuestos - Santana Hnos.

Genera archivos profesionales con formato, hojas múltiples y metadatos.
"""

import json
import os
import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
    numbers,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import PROXIMIDAD_ITUZAINGO, ZONAS


# ---------------------------------------------------------------------------
# Styling constants
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill(start_color="1A5C2A", end_color="1A5C2A", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

_ALT_ROW_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
_WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

_GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

_THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

_BODY_FONT = Font(name="Calibri", size=10)
_BODY_ALIGNMENT = Alignment(vertical="center", wrap_text=False)


# ---------------------------------------------------------------------------
# Helper: apply common sheet formatting
# ---------------------------------------------------------------------------

def _style_header_row(ws: Worksheet) -> None:
    """Apply header styling to the first row of a worksheet."""
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _THIN_BORDER


def _style_data_rows(ws: Worksheet, rating_col: Optional[int] = None) -> None:
    """
    Apply alternating row colors and body styling.
    If rating_col is given (1-based), apply conditional fill for rating values.
    """
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        fill = _ALT_ROW_FILL if row_idx % 2 == 0 else _WHITE_FILL
        for cell in row:
            cell.fill = fill
            cell.font = _BODY_FONT
            cell.alignment = _BODY_ALIGNMENT
            cell.border = _THIN_BORDER

        # Conditional formatting on rating
        if rating_col is not None and rating_col <= len(row):
            rating_cell = row[rating_col - 1]
            try:
                val = float(rating_cell.value) if rating_cell.value is not None else None
            except (ValueError, TypeError):
                val = None
            if val is not None:
                if val >= 4.0:
                    rating_cell.fill = _GREEN_FILL
                elif val < 3.0:
                    rating_cell.fill = _RED_FILL


def _auto_column_widths(ws: Worksheet, max_width: int = 45) -> None:
    """Auto-adjust column widths based on content (with a cap)."""
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws.max_row):
            for cell in row:
                if cell.value is not None:
                    cell_len = len(str(cell.value))
                    if cell_len > max_len:
                        max_len = cell_len
        adjusted = min(max_len + 3, max_width)
        ws.column_dimensions[col_letter].width = max(adjusted, 10)


def _set_phone_columns_as_text(ws: Worksheet, phone_col_indices: List[int]) -> None:
    """Format phone columns as text to prevent Excel from interpreting them as numbers."""
    for col_idx in phone_col_indices:
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.number_format = numbers.FORMAT_TEXT
                if cell.value is not None:
                    cell.value = str(cell.value)


def _add_autofilter(ws: Worksheet) -> None:
    """Add auto-filter to the header row."""
    if ws.max_row >= 1 and ws.max_column >= 1:
        last_col = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"


def _write_df_to_sheet(
    ws: Worksheet,
    df: pd.DataFrame,
    rating_col_name: Optional[str] = None,
    phone_col_names: Optional[List[str]] = None,
) -> None:
    """Write a DataFrame to a worksheet with full styling."""
    # Write headers
    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    # Write data
    for row_idx, row_data in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if pd.isna(value):
                cell.value = ""
            else:
                cell.value = value

    # Style
    _style_header_row(ws)

    # Find rating column index (1-based)
    rating_col_idx = None
    if rating_col_name and rating_col_name in df.columns:
        rating_col_idx = list(df.columns).index(rating_col_name) + 1

    _style_data_rows(ws, rating_col=rating_col_idx)

    # Phone columns as text
    if phone_col_names:
        phone_indices = []
        for pname in phone_col_names:
            if pname in df.columns:
                phone_indices.append(list(df.columns).index(pname) + 1)
        if phone_indices:
            _set_phone_columns_as_text(ws, phone_indices)

    _auto_column_widths(ws)
    _add_autofilter(ws)
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Column definitions per sheet
# ---------------------------------------------------------------------------

# All possible columns for "Base Completa"
_ALL_COLUMNS = [
    "nombre", "categoria_google", "direccion", "localidad", "partido", "zona",
    "telefono", "celular_whatsapp", "es_celular", "sitio_web",
    "instagram", "facebook", "email",
    "rating", "cantidad_resenas", "horarios",
    "google_maps_url", "notas",
]


def _safe_df(records: List[Dict[str, Any]], columns: List[str]) -> pd.DataFrame:
    """Create a DataFrame ensuring all specified columns exist."""
    df = pd.DataFrame(records)
    for col in columns:
        if col not in df.columns:
            df[col] = ""
    return df[columns]


def _display_columns(columns: List[str]) -> List[str]:
    """Convert internal column names to display-friendly Spanish headers."""
    mapping = {
        "nombre": "Nombre",
        "categoria_google": "Categoría",
        "direccion": "Dirección",
        "localidad": "Localidad",
        "partido": "Partido",
        "zona": "Zona",
        "telefono": "Teléfono",
        "celular_whatsapp": "WhatsApp",
        "es_celular": "Es Celular",
        "sitio_web": "Sitio Web",
        "instagram": "Instagram",
        "facebook": "Facebook",
        "email": "Email",
        "rating": "Rating",
        "cantidad_resenas": "Reseñas",
        "horarios": "Horarios",
        "google_maps_url": "Google Maps URL",
        "notas": "Notas",
        "distancia_estimada": "Distancia Est.",
    }
    return [mapping.get(c, c) for c in columns]


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def export_excel(records: List[Dict[str, Any]], filepath: str) -> str:
    """
    Create a professional Excel workbook with 6 sheets.

    Sheets:
        1. Base Completa      - All records, all columns
        2. Por Zona           - Summary pivot by zona/partido
        3. Con WhatsApp       - Records with celular/whatsapp
        4. Con Sitio Web      - Records with website
        5. Top Rating         - All records sorted by rating desc
        6. Zona Oeste Prioridad - Zona Oeste sorted by proximity to Ituzaingó

    Returns the filepath of the created workbook.
    """
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # ---- Sheet 1: Base Completa ----
    ws1 = wb.active
    ws1.title = "Base Completa"

    columns_all = [c for c in _ALL_COLUMNS if c in _get_available_cols(records)]
    # Ensure all desired columns are present
    for c in _ALL_COLUMNS:
        if c not in columns_all:
            columns_all.append(c)

    df_all = _safe_df(records, columns_all)
    df_all.columns = _display_columns(columns_all)

    _write_df_to_sheet(
        ws1, df_all,
        rating_col_name="Rating",
        phone_col_names=["Teléfono", "WhatsApp"],
    )

    # ---- Sheet 2: Por Zona ----
    ws2 = wb.create_sheet("Por Zona")
    zona_data = _build_zona_summary(records)
    df_zona = pd.DataFrame(zona_data, columns=["Zona", "Partido", "Cantidad", "% del Total"])
    _write_df_to_sheet(ws2, df_zona)

    # ---- Sheet 3: Con WhatsApp ----
    ws3 = wb.create_sheet("Con WhatsApp")
    whatsapp_cols = ["nombre", "localidad", "partido", "celular_whatsapp", "telefono", "rating", "google_maps_url"]
    whatsapp_records = [r for r in records if (r.get("celular_whatsapp") or "").strip()]
    df_wa = _safe_df(whatsapp_records, whatsapp_cols)
    df_wa.columns = _display_columns(whatsapp_cols)
    _write_df_to_sheet(
        ws3, df_wa,
        rating_col_name="Rating",
        phone_col_names=["WhatsApp", "Teléfono"],
    )

    # ---- Sheet 4: Con Sitio Web ----
    ws4 = wb.create_sheet("Con Sitio Web")
    web_cols = ["nombre", "localidad", "sitio_web", "instagram", "facebook", "rating"]
    web_records = [r for r in records if (r.get("sitio_web") or "").strip()]
    df_web = _safe_df(web_records, web_cols)
    df_web.columns = _display_columns(web_cols)
    _write_df_to_sheet(ws4, df_web, rating_col_name="Rating")

    # ---- Sheet 5: Top Rating ----
    ws5 = wb.create_sheet("Top Rating")
    rating_cols = ["nombre", "localidad", "rating", "cantidad_resenas", "categoria_google", "telefono", "direccion"]
    df_rating = _safe_df(records, rating_cols)
    df_rating.columns = _display_columns(rating_cols)

    # Sort by rating descending (handle non-numeric)
    df_rating["_sort_rating"] = pd.to_numeric(df_rating["Rating"], errors="coerce").fillna(0)
    df_rating = df_rating.sort_values("_sort_rating", ascending=False).drop(columns=["_sort_rating"])
    df_rating = df_rating.reset_index(drop=True)

    _write_df_to_sheet(
        ws5, df_rating,
        rating_col_name="Rating",
        phone_col_names=["Teléfono"],
    )

    # ---- Sheet 6: Zona Oeste Prioridad ----
    ws6 = wb.create_sheet("Zona Oeste Prioridad")
    oeste_cols = ["nombre", "localidad", "distancia_estimada", "telefono", "celular_whatsapp", "rating", "direccion"]
    oeste_records = _build_zona_oeste_priority(records)
    df_oeste = _safe_df(oeste_records, oeste_cols)
    df_oeste.columns = _display_columns(oeste_cols)
    _write_df_to_sheet(
        ws6, df_oeste,
        rating_col_name="Rating",
        phone_col_names=["Teléfono", "WhatsApp"],
    )

    # Save
    wb.save(str(path))
    return str(path)


def _get_available_cols(records: List[Dict[str, Any]]) -> set:
    """Get all column names present in at least one record."""
    cols = set()
    for r in records:
        cols.update(r.keys())
    return cols


def _build_zona_summary(records: List[Dict[str, Any]]) -> List[List[Any]]:
    """Build zona/partido summary rows sorted by count descending."""
    from collections import Counter

    total = len(records) if records else 1
    counts: Counter = Counter()

    for r in records:
        zona = r.get("zona") or "Sin Zona"
        partido = r.get("partido") or "Sin Partido"
        counts[(zona, partido)] += 1

    rows = []
    for (zona, partido), count in counts.most_common():
        pct = round((count / total) * 100, 1)
        rows.append([zona, partido, count, f"{pct}%"])

    return rows


def _build_zona_oeste_priority(records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Filter Zona Oeste records and sort by proximity to Ituzaingó.
    Adds a 'distancia_estimada' field based on position in PROXIMIDAD_ITUZAINGO.
    """
    oeste_records = [r for r in records if (r.get("zona") or "").strip() == "Zona Oeste"]

    # Build proximity index
    prox_index = {loc: idx for idx, loc in enumerate(PROXIMIDAD_ITUZAINGO)}
    max_pos = len(PROXIMIDAD_ITUZAINGO)

    for rec in oeste_records:
        localidad = (rec.get("localidad") or "").strip()
        pos = prox_index.get(localidad, max_pos)
        rec["distancia_estimada"] = pos + 1  # 1-based ranking

    # Sort by proximity
    oeste_records.sort(key=lambda r: r.get("distancia_estimada", max_pos + 1))

    return oeste_records


# ---------------------------------------------------------------------------
# CSV export
# ---------------------------------------------------------------------------

def export_csv(records: List[Dict[str, Any]], filepath: str) -> str:
    """
    Export records to CSV with UTF-8 BOM and semicolon separator.

    Returns the filepath of the created file.
    """
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    # Build DataFrame with all columns
    available_cols = []
    for c in _ALL_COLUMNS:
        available_cols.append(c)

    df = _safe_df(records, available_cols)
    df.columns = _display_columns(available_cols)

    df.to_csv(
        str(path),
        sep=";",
        index=False,
        encoding="utf-8-sig",  # UTF-8 with BOM
    )

    return str(path)


# ---------------------------------------------------------------------------
# JSON export
# ---------------------------------------------------------------------------

def export_json(records: List[Dict[str, Any]], filepath: str) -> str:
    """
    Export records to structured JSON with metadata and nested format.

    Returns the filepath of the created file.
    """
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    # Collect zonas covered
    zonas_cubiertas = sorted({
        r.get("zona") or "Sin Zona" for r in records
    })

    # Count unique queries (if tracked)
    queries_count = len({r.get("query") for r in records if r.get("query")})
    if queries_count == 0:
        queries_count = 1  # At least 1 if we have data

    metadata = {
        "fecha_scraping": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_registros": len(records),
        "zonas_cubiertas": zonas_cubiertas,
        "queries_ejecutados": queries_count,
    }

    clientes = []
    for rec in records:
        # Parse coordinates if available
        lat = 0.0
        lng = 0.0
        coords = rec.get("coordenadas")
        if isinstance(coords, dict):
            lat = coords.get("lat", 0.0)
            lng = coords.get("lng", 0.0)
        elif isinstance(coords, str) and "," in coords:
            try:
                parts = coords.split(",")
                lat = float(parts[0].strip())
                lng = float(parts[1].strip())
            except (ValueError, IndexError):
                pass

        # Parse rating
        rating_val = 0.0
        try:
            rating_val = float(rec.get("rating") or 0)
        except (ValueError, TypeError):
            pass

        # Parse review count
        resenas_val = 0
        try:
            resenas_raw = str(rec.get("cantidad_resenas") or "0")
            resenas_clean = re.sub(r"[^\d]", "", resenas_raw)
            resenas_val = int(resenas_clean) if resenas_clean else 0
        except (ValueError, TypeError):
            pass

        cliente = {
            "id": str(uuid.uuid4()),
            "nombre": rec.get("nombre") or "",
            "tipo": rec.get("categoria_google") or "",
            "contacto": {
                "telefono": rec.get("telefono") or "",
                "celular": rec.get("celular_whatsapp") or "",
                "email": rec.get("email") or "",
                "whatsapp": rec.get("celular_whatsapp") or "",
            },
            "ubicacion": {
                "direccion": rec.get("direccion") or "",
                "localidad": rec.get("localidad") or "",
                "partido": rec.get("partido") or "",
                "zona": rec.get("zona") or "",
                "coordenadas": {
                    "lat": lat,
                    "lng": lng,
                },
            },
            "online": {
                "sitio_web": rec.get("sitio_web") or "",
                "instagram": rec.get("instagram") or "",
                "facebook": rec.get("facebook") or "",
                "google_maps": rec.get("google_maps_url") or "",
            },
            "reputacion": {
                "rating": rating_val,
                "cantidad_resenas": resenas_val,
            },
            "estado_comercial": "nuevo",
            "notas": rec.get("notas") or "",
        }

        clientes.append(cliente)

    output = {
        "metadata": metadata,
        "clientes_potenciales": clientes,
    }

    with open(str(path), "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    return str(path)


# ---------------------------------------------------------------------------
# Export all
# ---------------------------------------------------------------------------

def export_all(
    records: List[Dict[str, Any]],
    output_dir: str = "data/output",
) -> List[str]:
    """
    Export records to XLSX, CSV, and JSON.

    Filenames include today's date: repuestos_amba_YYYYMMDD.{ext}

    Returns list of created filepaths. Prints summary to console.
    """
    out_path = Path(output_dir)
    out_path.mkdir(parents=True, exist_ok=True)

    date_str = datetime.now().strftime("%Y%m%d")
    base_name = f"repuestos_amba_{date_str}"

    xlsx_path = str(out_path / f"{base_name}.xlsx")
    csv_path = str(out_path / f"{base_name}.csv")
    json_path = str(out_path / f"{base_name}.json")

    created = []

    # Excel
    export_excel(records, xlsx_path)
    created.append(xlsx_path)

    # CSV
    export_csv(records, csv_path)
    created.append(csv_path)

    # JSON
    export_json(records, json_path)
    created.append(json_path)

    # Summary
    total = len(records)
    zonas = {}
    for r in records:
        z = r.get("zona") or "Sin Zona"
        zonas[z] = zonas.get(z, 0) + 1

    con_whatsapp = sum(1 for r in records if (r.get("celular_whatsapp") or "").strip())
    con_web = sum(1 for r in records if (r.get("sitio_web") or "").strip())

    print("\n" + "=" * 60)
    print("  EXPORTACIÓN COMPLETADA - Santana Hnos.")
    print("=" * 60)
    print(f"  Total registros: {total}")
    print(f"  Con WhatsApp:    {con_whatsapp}")
    print(f"  Con Sitio Web:   {con_web}")
    print()
    print("  Distribución por zona:")
    for zona, count in sorted(zonas.items(), key=lambda x: -x[1]):
        pct = round((count / max(total, 1)) * 100, 1)
        print(f"    {zona}: {count} ({pct}%)")
    print()
    print("  Archivos generados:")
    for fp in created:
        size_kb = os.path.getsize(fp) / 1024
        print(f"    {fp} ({size_kb:.1f} KB)")
    print("=" * 60 + "\n")

    return created
