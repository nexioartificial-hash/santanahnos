"""
Exportador de resultados a Excel, CSV y JSON.
Genera archivos formateados con la marca Santana Hnos.
"""

import json
import os
from datetime import datetime
from typing import List, Dict, Any, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


OUTPUT_DIR = os.path.join(os.path.dirname(__file__), 'data', 'output')

# Colores Santana Hnos.
GREEN_DARK = '1A5C2A'
GREEN_LIGHT = 'E8F5E9'
YELLOW_LIGHT = 'FFF9C4'
ORANGE_LIGHT = 'FFE0B2'
WHITE = 'FFFFFF'
GRAY_LIGHT = 'F5F5F5'


def _ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def _build_dataframe(results: List[dict], extra_columns: Optional[List[str]] = None) -> pd.DataFrame:
    """Construye un DataFrame con los resultados."""
    rows = []
    for r in results:
        row = {
            'Numero Original': r.get('original', ''),
            'Normalizado Local': r.get('normalizado_local', ''),
            'Normalizado Internacional': r.get('normalizado_internacional', ''),
            'Tipo': r.get('tipo', '').capitalize(),
            'Confianza': r.get('confianza', '').capitalize(),
            'Codigo de Area': r.get('codigo_area', ''),
            'WhatsApp Probable': 'Si' if r.get('whatsapp_probable') else 'No',
            'WhatsApp Verificado': {True: 'Si', False: 'No', None: 'No verificado'}.get(
                r.get('whatsapp_verificado'), 'No verificado'
            ),
            'Link wa.me': r.get('link_whatsapp', ''),
            'Notas': r.get('notas', ''),
        }

        # Agregar datos extra (nombre, direccion, etc.)
        extras = r.get('datos_extra', {})
        if extras:
            for key, val in extras.items():
                col_name = key.replace('_', ' ').title()
                row[col_name] = val

        rows.append(row)

    df = pd.DataFrame(rows)
    return df


def _style_worksheet(ws, df: pd.DataFrame):
    """Aplica estilos al worksheet de Excel."""
    header_font = Font(name='Calibri', bold=True, color=WHITE, size=11)
    header_fill = PatternFill(start_color=GREEN_DARK, end_color=GREEN_DARK, fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    celular_fill = PatternFill(start_color=GREEN_LIGHT, end_color=GREEN_LIGHT, fill_type='solid')
    no_clasif_fill = PatternFill(start_color=YELLOW_LIGHT, end_color=YELLOW_LIGHT, fill_type='solid')
    baja_fill = PatternFill(start_color=ORANGE_LIGHT, end_color=ORANGE_LIGHT, fill_type='solid')
    alt_fill = PatternFill(start_color=GRAY_LIGHT, end_color=GRAY_LIGHT, fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD'),
    )

    # Headers
    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Encontrar indices de columnas clave
    cols = list(df.columns)
    tipo_idx = cols.index('Tipo') + 1 if 'Tipo' in cols else None
    conf_idx = cols.index('Confianza') + 1 if 'Confianza' in cols else None

    # Datos
    for row_idx in range(2, len(df) + 2):
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

        # Color de fila segun tipo
        if tipo_idx:
            tipo_val = ws.cell(row=row_idx, column=tipo_idx).value
            if tipo_val == 'Celular':
                for col_idx in range(1, len(df.columns) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = celular_fill
            elif tipo_val == 'No_clasificado':
                for col_idx in range(1, len(df.columns) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = no_clasif_fill
            elif (row_idx - 2) % 2 == 1:
                for col_idx in range(1, len(df.columns) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = alt_fill

        # Resaltar confianza baja
        if conf_idx:
            conf_val = ws.cell(row=row_idx, column=conf_idx).value
            if conf_val == 'Baja':
                for col_idx in range(1, len(df.columns) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = baja_fill

    # Auto-ajustar anchos
    for col_idx in range(1, len(df.columns) + 1):
        max_len = len(str(ws.cell(row=1, column=col_idx).value or ''))
        for row_idx in range(2, min(len(df) + 2, 100)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 40)

    # Filtros automaticos
    ws.auto_filter.ref = ws.dimensions


def export_excel(results: List[dict], stats: dict, filename: str = None) -> str:
    """Exporta resultados a Excel con multiples hojas."""
    _ensure_output_dir()
    if not filename:
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'telefonos_analisis_{ts}.xlsx'
    filepath = os.path.join(OUTPUT_DIR, filename)

    df_all = _build_dataframe(results)
    df_cel = df_all[df_all['Tipo'] == 'Celular'].copy()
    df_fijo = df_all[df_all['Tipo'] == 'Fijo'].copy()

    # Hoja de estadisticas
    stats_data = {
        'Metrica': [
            'Total numeros', 'Lineas fijas', 'Celulares', 'No clasificados',
            '', 'WhatsApp probable', 'WhatsApp verificado',
            '', 'Confianza Alta', 'Confianza Media', 'Confianza Baja',
        ],
        'Cantidad': [
            stats['total'], stats['fijos'], stats['celulares'], stats['no_clasificados'],
            '', stats['whatsapp_probable'], stats['whatsapp_verificado'],
            '', stats['confianza_alta'], stats['confianza_media'], stats['confianza_baja'],
        ],
        'Porcentaje': [
            '100%', f"{stats['pct_fijos']}%", f"{stats['pct_celulares']}%", f"{stats['pct_no_clasificados']}%",
            '', '', '',
            '', f"{stats['pct_confianza_alta']}%", f"{stats['pct_confianza_media']}%", f"{stats['pct_confianza_baja']}%",
        ],
    }
    df_stats = pd.DataFrame(stats_data)

    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df_all.to_excel(writer, sheet_name='Resultados', index=False)
        df_cel.to_excel(writer, sheet_name='Solo Celulares', index=False)
        df_fijo.to_excel(writer, sheet_name='Solo Fijos', index=False)
        df_stats.to_excel(writer, sheet_name='Estadisticas', index=False)

    # Aplicar estilos
    wb = load_workbook(filepath)
    for sheet_name in ['Resultados', 'Solo Celulares', 'Solo Fijos']:
        ws = wb[sheet_name]
        df_sheet = {'Resultados': df_all, 'Solo Celulares': df_cel, 'Solo Fijos': df_fijo}[sheet_name]
        if not df_sheet.empty:
            _style_worksheet(ws, df_sheet)

    # Estilo basico para estadisticas
    ws_stats = wb['Estadisticas']
    for col_idx in range(1, 4):
        cell = ws_stats.cell(row=1, column=col_idx)
        cell.font = Font(name='Calibri', bold=True, color=WHITE, size=11)
        cell.fill = PatternFill(start_color=GREEN_DARK, end_color=GREEN_DARK, fill_type='solid')

    wb.save(filepath)
    return filepath


def export_csv(results: List[dict], filename: str = None) -> str:
    """Exporta resultados a CSV con separador ; y UTF-8 BOM."""
    _ensure_output_dir()
    if not filename:
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'telefonos_analisis_{ts}.csv'
    filepath = os.path.join(OUTPUT_DIR, filename)

    df = _build_dataframe(results)
    df.to_csv(filepath, index=False, sep=';', encoding='utf-8-sig')
    return filepath


def export_json(results: List[dict], stats: dict, filename: str = None) -> str:
    """Exporta resultados a JSON compatible con el panel admin."""
    _ensure_output_dir()
    if not filename:
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'telefonos_analisis_{ts}.json'
    filepath = os.path.join(OUTPUT_DIR, filename)

    output = {
        'metadata': {
            'fecha_analisis': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'total_numeros': stats['total'],
            'fijos': stats['fijos'],
            'celulares': stats['celulares'],
            'no_clasificados': stats['no_clasificados'],
            'whatsapp_verificados': stats['whatsapp_verificado'],
        },
        'numeros': [
            {
                'original': r.get('original', ''),
                'normalizado_local': r.get('normalizado_local', ''),
                'normalizado_internacional': r.get('normalizado_internacional', ''),
                'tipo': r.get('tipo', ''),
                'confianza': r.get('confianza', ''),
                'codigo_area': r.get('codigo_area', ''),
                'whatsapp_probable': r.get('whatsapp_probable', False),
                'whatsapp_verificado': r.get('whatsapp_verificado'),
                'link_whatsapp': r.get('link_whatsapp'),
                'notas': r.get('notas', ''),
                'datos_extra': r.get('datos_extra', {}),
            }
            for r in results
        ],
    }

    with open(filepath, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2)
    return filepath


def export_all(results: List[dict], stats: dict) -> dict:
    """Exporta en todos los formatos y retorna las rutas."""
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    paths = {
        'excel': export_excel(results, stats, f'telefonos_{ts}.xlsx'),
        'csv': export_csv(results, f'telefonos_{ts}.csv'),
        'json': export_json(results, stats, f'telefonos_{ts}.json'),
    }
    return paths
